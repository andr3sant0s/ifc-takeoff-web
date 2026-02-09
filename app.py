
from flask import Flask, request, render_template_string, send_file, jsonify
import os, gc
import pandas as pd
import ifcopenshell
import ifcopenshell.util.element as Element
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 300 * 1024 * 1024

UPLOAD_FOLDER = "uploads"
OUTPUT_FOLDER = "outputs"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

progress_state = {"current":0, "total":1, "text":"idle"}

def set_progress(c,t,txt):
    progress_state["current"]=c
    progress_state["total"]=t
    progress_state["text"]=txt

HTML = "<h2>IFC Takeoff Free Safe</h2>"

def txt(v): 
    return str(v).strip().lower() if v else ""

def get_pset_value(el,ps,p):
    try:
        psets=Element.get_psets(el)
        if psets and ps in psets:
            return psets[ps].get(p)
    except: pass
    return None

def get_quantity(el,uom):
    uom=txt(uom)
    if uom=="un": return 1

    if uom=="m":
        L=get_pset_value(el,"BaseQuantities","Length")
        if L: return float(L)

    if uom=="m3":
        v=get_pset_value(el,"BaseQuantities","NetVolume")
        if v: return float(v)

    return 0

def run_one(ifc_path, excel_path, output_path):
    model=ifcopenshell.open(ifc_path)
    elements=model.by_type("IfcElement")

    df=pd.read_excel(excel_path,dtype=str)

    out=[]; val=[]
    total_el=len(elements)
    step=0

    for _,row in df.iterrows():
        wbs=txt(row["WBS"])
        desc=row["Description"]
        uom=row["UN"]
        cost=float(str(row["Cost"]).replace(",", "."))

        by={}

        for el in elements:
            step+=1
            set_progress(step,total_el,os.path.basename(ifc_path))

            if step % 200 == 0:
                gc.collect()

            el_wbs=txt(get_pset_value(el,"Pset_+351","WBS"))
            piso=get_pset_value(el,"Pset_+351","Piso")

            if not piso:
                val.append({"GlobalId":el.GlobalId,"IfcClass":el.is_a(),"WBS":el_wbs,"Reason":"Missing Piso"})
                piso="NO PISO"

            if el_wbs!=wbs: continue

            q=get_quantity(el,uom)
            by[piso]=by.get(piso,0)+q

        tq=round(sum(by.values()),3)
        tc=round(tq*cost,3)

        out.append([row["WBS"],desc+" - TOTAL",uom,"",tq,cost,tc])

    wb=load_workbook(excel_path)
    ws=wb.active
    ws.delete_rows(2,ws.max_row)

    for r,data in enumerate(out,2):
        for c,v in enumerate(data,1):
            ws.cell(row=r,column=c).value=v

    wb.save(output_path)

@app.route('/',methods=['GET','POST'])
def index():
    if request.method=='POST':
        excel=request.files['excel']
        excel_path=os.path.join(UPLOAD_FOLDER,secure_filename(excel.filename))
        excel.save(excel_path)

        ifc=request.files['ifc']
        p=os.path.join(UPLOAD_FOLDER,secure_filename(ifc.filename))
        ifc.save(p)

        o=os.path.join(OUTPUT_FOLDER,"result.xlsx")
        run_one(p,excel_path,o)

        return send_file(o,as_attachment=True)

    return render_template_string(HTML)

if __name__=='__main__':
    app.run(host="0.0.0.0", port=10000)
